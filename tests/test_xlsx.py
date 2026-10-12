"""XLSX export: openpyxl can open it, totals match the DB."""
from __future__ import annotations

import datetime

from openpyxl import load_workbook
from sqlmodel import Session

from app.auth.passwords import hash_password
from app.db import get_engine
from app.models.account import Account
from app.models.category import Category
from app.models.transaction import Transaction
from app.models.user import User
from app.routers.auth import _seed_default_categories
from app.services import exports as exports_service


def _setup() -> tuple[int, int]:
    with Session(get_engine()) as s:
        u = User(
            username="alice",
            password_hash=hash_password("pw12345678"),
            base_currency="USD",
            timezone="UTC",
        )
        s.add(u)
        s.commit()
        s.refresh(u)
        _seed_default_categories(s, u.id)
        a = Account(user_id=u.id, name="Checking", type="bank")
        s.add(a)
        s.commit()
        s.refresh(a)
        groceries = (
            s.query(Category).filter(Category.user_id == u.id, Category.name == "Groceries").first()
        )
        s.add(
            Transaction(
                user_id=u.id,
                account_id=a.id,
                date=datetime.date(2026, 10, 1),
                amount_minor=-1000,
                currency="USD",
                description="x",
                category_id=groceries.id,
                dedup_key="k1",
            )
        )
        s.add(
            Transaction(
                user_id=u.id,
                account_id=a.id,
                date=datetime.date(2026, 10, 2),
                amount_minor=5000,
                currency="USD",
                description="salary",
                dedup_key="k2",
            )
        )
        s.add(
            Transaction(
                user_id=u.id,
                account_id=a.id,
                date=datetime.date(2026, 10, 3),
                amount_minor=-2000,
                currency="USD",
                description="xfer",
                is_transfer=True,
                dedup_key="k3",
            )
        )
        s.commit()
        return u.id, groceries.id


def test_xlsx_opens_and_totals_match() -> None:
    uid, groceries = _setup()
    with Session(get_engine()) as s:
        body = exports_service.export_xlsx(s, uid)
    assert isinstance(body, bytes)
    wb = load_workbook(__import__("io").BytesIO(body))
    assert "Transactions" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    # The Summary sheet ends with Net = income - expense = 50 - 10 = 40.
    summary = wb["Summary"]
    # Find the row with "Net".
    net = None
    for row in summary.iter_rows(values_only=True):
        if row and row[0] == "Net":
            net = row[1]
            break
    assert net is not None
    assert abs(float(net) - 40.0) < 1e-6
    # The Transactions sheet has 3 rows + header.
    ws = wb["Transactions"]
    assert ws.max_row == 4


def test_xlsx_transfers_excluded_from_totals() -> None:
    uid, _ = _setup()
    with Session(get_engine()) as s:
        body = exports_service.export_xlsx(s, uid)
    wb = load_workbook(__import__("io").BytesIO(body))
    summary = wb["Summary"]
    rows = {r[0]: r[1] for r in summary.iter_rows(values_only=True) if r and r[0]}
    # The transfer row (-2000) is excluded; income=50, expense=10, net=40.
    assert abs(float(rows["Income"]) - 50.0) < 1e-6
    assert abs(float(rows["Expense"]) - 10.0) < 1e-6
    assert abs(float(rows["Net"]) - 40.0) < 1e-6
